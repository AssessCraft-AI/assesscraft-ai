"""Course material categories: lecture slides (incl. .pptx) and model papers."""
import io

import pytest
from httpx import AsyncClient, ASGITransport
from pptx import Presentation

from app.main import app
from seed.load_seed import seed_database

LECTURER = {"X-Demo-Role": "lecturer"}


@pytest.fixture(autouse=True)
async def setup():
    await seed_database()


def _deck() -> bytes:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "B+ Tree Indexing"
    slide.placeholders[1].text = "Leaf nodes are linked"
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_upload_slides_and_model_paper():
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        pptx_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        slides = await c.post(
            "/courses/CS2040/documents",
            headers=LECTURER,
            files={"file": ("week5.pptx", _deck(), pptx_type)},
            data={"kind": "slides"},
        )
        assert slides.status_code == 201, slides.text
        assert slides.json()["kind"] == "slides"
        assert "B+ Tree Indexing" in slides.json()["contentSnippet"]

        model = await c.post(
            "/courses/CS2040/documents",
            headers=LECTURER,
            files={"file": ("model.md", b"Q1. Define a transaction.", "text/markdown")},
            data={"kind": "model_paper", "year": "2026"},
        )
        assert model.status_code == 201, model.text
        assert model.json()["kind"] == "model_paper" and model.json()["year"] == 2026

        bad = await c.post(
            "/courses/CS2040/documents",
            headers=LECTURER,
            files={"file": ("x.md", b"x", "text/markdown")},
            data={"kind": "brochure"},
        )
        assert bad.status_code == 422
